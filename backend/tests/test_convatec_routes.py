import io

import openpyxl
import pandas as pd
import pytest

from app import create_app

# Convatec's blueprint is disconnected from create_app() on request (see
# app/__init__.py) -- these routes no longer exist, so every test below would
# just fail on a 404 rather than exercise anything real. Skipped, not
# deleted: the module's code/tests stay on disk in case the demo work
# resumes.
pytestmark = pytest.mark.skip(reason="Convatec module is disconnected from the app (blueprint not registered)")


@pytest.fixture
def client():
    app = create_app()
    return app.test_client()


def _maestros_xlsx_bytes():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Tables"
    ws["B2"] = "PRODUCTO"
    ws["K2"] = "REPRESENTANTES"
    ws["O2"] = "CONVENIOS"

    # Joined by ICC (HyperSoft's own code), not SAP Code -- see
    # pipeline.homologar_productos docstring for why.
    for i, h in enumerate(["ICC", "Franquicia", "Familia", "Clase", "Grupo", "Clasificación Cuota Producto"]):
        ws.cell(row=4, column=2 + i, value=h)
    ws.cell(row=5, column=2, value="1004113")
    ws.cell(row=5, column=3, value="CONTINENCE")
    ws.cell(row=5, column=4, value="Flexiseal")
    ws.cell(row=5, column=5, value="0")
    ws.cell(row=5, column=6, value="Continence Care")
    ws.cell(row=5, column=7, value="1. CCC")

    for i, h in enumerate(["Convenio", "Grupo de vendedores", "Representante"]):
        ws.cell(row=4, column=11 + i, value=h)
    ws.cell(row=5, column=11, value="SALUD TOTAL")
    ws.cell(row=5, column=12, value=45)
    ws.cell(row=5, column=13, value="Katerine Bonilla")

    for i, h in enumerate(["CONVENIO", "CIUDAD", "CLIENTE", "CODIGO SAP"]):
        ws.cell(row=4, column=15 + i, value=h)
    ws.cell(row=5, column=15, value="SALUD TOTAL")
    ws.cell(row=5, column=16, value="BOGOTA")

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


def _productos_xlsx_bytes():
    buf = io.BytesIO()
    rows = [
        {
            "Codigo": "1004113",
            "Convenio": "X",
            "Departamento": "META",
            "Valor Total": 100.0,
            "Paciente": "Jose Perez",
            "Identificación": "123",
        },
        {
            "Codigo": "1004113",
            "Convenio": "X",
            "Departamento": "ANTIOQUIA",
            "Valor Total": 200.0,
            "Paciente": "Ana Ruiz",
            "Identificación": "456",
        },
    ]
    pd.DataFrame(rows).to_excel(buf, index=False)
    buf.seek(0)
    return buf


def _envios_nacionales_xlsx_bytes():
    buf = io.BytesIO()
    rows = [{"Departamento": "META", "Valor Total": 30.0, "Paciente": "Filtro Inesperado"}]
    pd.DataFrame(rows).to_excel(buf, index=False)
    buf.seek(0)
    return buf


HEADERS = {"X-Session-Id": "test-session"}


def test_full_cycle_upload_maestros_ventas_procesar_and_descargar(client):
    maestros_resp = client.post(
        "/api/convatec/tablas-maestras",
        data={"file": (_maestros_xlsx_bytes(), "maestros.xlsx")},
        content_type="multipart/form-data",
        headers=HEADERS,
    )
    assert maestros_resp.status_code == 200
    assert maestros_resp.get_json()["success"] is True

    ventas_resp = client.post(
        "/api/convatec/ventas/productos",
        data={"file": (_productos_xlsx_bytes(), "productos.xlsx")},
        content_type="multipart/form-data",
        headers=HEADERS,
    )
    assert ventas_resp.status_code == 200
    ventas_body = ventas_resp.get_json()["data"]
    # PHI columns never make it past ingestion, even for a session-scoped upload.
    assert "Paciente" not in ventas_body["columns"]
    assert "Identificación" not in ventas_body["columns"]

    procesar_resp = client.post(
        "/api/convatec/procesar", json={"mes": 3}, headers=HEADERS
    )
    assert procesar_resp.status_code == 200
    procesar_body = procesar_resp.get_json()["data"]
    assert procesar_body["totalLineas"] == 2

    comision_resp = client.post(
        "/api/convatec/comision-ilustrativa",
        json={"valorReferencia": 100},
        headers=HEADERS,
    )
    assert comision_resp.status_code == 200
    assert comision_resp.get_json()["data"]["esOficial"] is False

    preview_resp = client.get("/api/convatec/resultado-preview?limit=1", headers=HEADERS)
    assert preview_resp.status_code == 200
    preview_body = preview_resp.get_json()["data"]
    assert preview_body["totalRows"] == 2
    assert preview_body["previewRows"] == 1
    assert "Comision_Ilustrativa (VALOR DE EJEMPLO - NO OFICIAL)" in preview_body["columns"]

    descargar_resp = client.get("/api/convatec/descargar", headers=HEADERS)
    assert descargar_resp.status_code == 200
    assert descargar_resp.mimetype == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def test_envios_nacionales_upload_also_drops_phi_columns_defensively(client):
    """NFR-2 must not rely on the assumption that envios-nacionales never
    carries patient data -- drop_patient_columns runs for every tipo."""
    resp = client.post(
        "/api/convatec/ventas/envios-nacionales",
        data={"file": (_envios_nacionales_xlsx_bytes(), "envios.xlsx")},
        content_type="multipart/form-data",
        headers=HEADERS,
    )
    assert resp.status_code == 200
    assert "Paciente" not in resp.get_json()["data"]["columns"]


def test_procesar_without_master_tables_fails_clearly(client):
    resp = client.post("/api/convatec/procesar", json={}, headers={"X-Session-Id": "sin-maestros"})
    assert resp.status_code == 400
    assert "tablas maestras" in resp.get_json()["message"]


def test_procesar_rejects_bool_as_mes(client):
    client.post(
        "/api/convatec/tablas-maestras",
        data={"file": (_maestros_xlsx_bytes(), "maestros.xlsx")},
        content_type="multipart/form-data",
        headers={"X-Session-Id": "bool-mes"},
    )
    resp = client.post(
        "/api/convatec/procesar", json={"mes": True}, headers={"X-Session-Id": "bool-mes"}
    )
    assert resp.status_code == 400


def test_comision_ilustrativa_rejects_bool_and_non_finite_valor(client):
    session = {"X-Session-Id": "comision-validacion"}
    client.post(
        "/api/convatec/tablas-maestras",
        data={"file": (_maestros_xlsx_bytes(), "maestros.xlsx")},
        content_type="multipart/form-data",
        headers=session,
    )
    client.post(
        "/api/convatec/ventas/productos",
        data={"file": (_productos_xlsx_bytes(), "productos.xlsx")},
        content_type="multipart/form-data",
        headers=session,
    )
    client.post("/api/convatec/procesar", json={"mes": 3}, headers=session)

    bool_resp = client.post("/api/convatec/comision-ilustrativa", json={"valorReferencia": True}, headers=session)
    assert bool_resp.status_code == 400

    inf_resp = client.post("/api/convatec/comision-ilustrativa", json={"valorReferencia": float("inf")}, headers=session)
    assert inf_resp.status_code == 400


def test_descargar_without_ciclo_procesado_fails_clearly(client):
    resp = client.get("/api/convatec/descargar", headers={"X-Session-Id": "sin-ciclo"})
    assert resp.status_code == 400


def test_preview_without_ciclo_procesado_fails_clearly(client):
    resp = client.get("/api/convatec/resultado-preview", headers={"X-Session-Id": "sin-ciclo-preview"})
    assert resp.status_code == 400


def test_preview_rejects_invalid_limit(client):
    session = {"X-Session-Id": "preview-limit"}
    client.post(
        "/api/convatec/tablas-maestras",
        data={"file": (_maestros_xlsx_bytes(), "maestros.xlsx")},
        content_type="multipart/form-data",
        headers=session,
    )
    client.post(
        "/api/convatec/ventas/productos",
        data={"file": (_productos_xlsx_bytes(), "productos.xlsx")},
        content_type="multipart/form-data",
        headers=session,
    )
    client.post("/api/convatec/procesar", json={"mes": 3}, headers=session)

    resp = client.get("/api/convatec/resultado-preview?limit=0", headers=session)
    assert resp.status_code == 400
    resp2 = client.get("/api/convatec/resultado-preview?limit=abc", headers=session)
    assert resp2.status_code == 400


def _reconocimiento_xlsx_bytes(total):
    import openpyxl

    wb = openpyxl.Workbook()
    ws = wb.active
    for i, h in enumerate(["JE line", "GL account", "Debit amount", "Credit amount", "PK"]):
        ws.cell(row=9, column=1 + i, value=h)
    ws.cell(row=10, column=1, value=1)
    ws.cell(row=10, column=4, value=total)
    ws.cell(row=10, column=5, value=50)
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


def test_validacion_cifras_matches_when_totals_are_equal(client):
    session = {"X-Session-Id": "cifras-cuadran"}
    client.post(
        "/api/convatec/tablas-maestras",
        data={"file": (_maestros_xlsx_bytes(), "maestros.xlsx")},
        content_type="multipart/form-data",
        headers=session,
    )
    client.post(
        "/api/convatec/ventas/productos",
        data={"file": (_productos_xlsx_bytes(), "productos.xlsx")},
        content_type="multipart/form-data",
        headers=session,
    )
    client.post("/api/convatec/procesar", json={"mes": 3}, headers=session)
    # _productos_xlsx_bytes() rows sum to 100.0 + 200.0 = 300.0
    client.post(
        "/api/convatec/reconocimiento-ingreso",
        data={"file": (_reconocimiento_xlsx_bytes(300.0), "reconocimiento.xlsx")},
        content_type="multipart/form-data",
        headers=session,
    )

    resp = client.get("/api/convatec/validacion-cifras", headers=session)
    body = resp.get_json()["data"]
    assert body["cifrasCuadran"] is True
    assert body["totalProcesado"] == 300.0
    assert body["totalReconocimiento"] == 300.0


def test_validacion_cifras_flags_mismatch(client):
    session = {"X-Session-Id": "cifras-no-cuadran"}
    client.post(
        "/api/convatec/tablas-maestras",
        data={"file": (_maestros_xlsx_bytes(), "maestros.xlsx")},
        content_type="multipart/form-data",
        headers=session,
    )
    client.post(
        "/api/convatec/ventas/productos",
        data={"file": (_productos_xlsx_bytes(), "productos.xlsx")},
        content_type="multipart/form-data",
        headers=session,
    )
    client.post("/api/convatec/procesar", json={"mes": 3}, headers=session)
    client.post(
        "/api/convatec/reconocimiento-ingreso",
        data={"file": (_reconocimiento_xlsx_bytes(999.0), "reconocimiento.xlsx")},
        content_type="multipart/form-data",
        headers=session,
    )

    resp = client.get("/api/convatec/validacion-cifras", headers=session)
    body = resp.get_json()["data"]
    assert body["cifrasCuadran"] is False
    assert body["diferencia"] != 0


def test_validacion_cifras_requires_both_inputs(client):
    resp = client.get("/api/convatec/validacion-cifras", headers={"X-Session-Id": "cifras-vacio"})
    assert resp.status_code == 400


def test_upload_ventas_rejects_unknown_tipo(client):
    resp = client.post(
        "/api/convatec/ventas/algo-invalido",
        data={"file": (_productos_xlsx_bytes(), "x.xlsx")},
        content_type="multipart/form-data",
        headers=HEADERS,
    )
    assert resp.status_code == 400


def _comision_externa_xlsx_bytes():
    import openpyxl

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Hoja2"
    for i, h in enumerate(["Vendedor", "Tipo", "No.Factura", "Dia", "Bod", "Gravado", "Exento", "Bruto", "I.V.A.", "Flete", "N E T O", "Com", "Comision"]):
        ws.cell(row=7, column=1 + i, value=h)
    ws.cell(row=8, column=1, value="VEGA BELEÑO CARLOS")
    ws.cell(row=8, column=11, value=100.0)
    ws.cell(row=8, column=12, value=20.0)
    ws.cell(row=8, column=13, value=20.0)
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


def test_comision_externa_referencia_is_stateless_and_never_touches_convatec_data(client):
    resp = client.post(
        "/api/convatec/comision-externa-referencia",
        data={"file": (_comision_externa_xlsx_bytes(), "comision.xls")},
        content_type="multipart/form-data",
        headers={"X-Session-Id": "referencia-externa"},
    )
    assert resp.status_code == 200
    filas = resp.get_json()["data"]["filas"]
    assert filas == [{"vendedor": "VEGA BELEÑO CARLOS", "neto": 100.0, "porcentaje": 20.0, "comision": 20.0}]

    # Never stored -- procesar/descargar for this session are unaffected.
    resp_procesar = client.post(
        "/api/convatec/procesar", json={}, headers={"X-Session-Id": "referencia-externa"}
    )
    assert resp_procesar.status_code == 400


def test_maestros_resumen_available_without_upload(client):
    resp = client.get("/api/convatec/maestros-resumen", headers={"X-Session-Id": "resumen-sin-upload"})
    assert resp.status_code == 200
    body = resp.get_json()["data"]
    assert len(body["continence"]) > 0
    assert len(body["enviosNacionales"]) > 0
    assert len(body["repartir"]) > 0
    assert body["otrasFranquiciasDisponible"] is False
    assert body["otrasFranquicias"] == []


def test_maestros_resumen_includes_otras_franquicias_after_upload(client):
    session = {"X-Session-Id": "resumen-con-upload"}
    client.post(
        "/api/convatec/tablas-maestras",
        data={"file": (_maestros_xlsx_bytes(), "maestros.xlsx")},
        content_type="multipart/form-data",
        headers=session,
    )
    resp = client.get("/api/convatec/maestros-resumen", headers=session)
    body = resp.get_json()["data"]
    assert body["otrasFranquiciasDisponible"] is True
    assert body["otrasFranquicias"][0]["convenio"] == "SALUD TOTAL"
    assert body["otrasFranquicias"][0]["ciudad"] == "BOGOTA"


def test_reiniciar_clears_ventas_but_keeps_global_maestros(client):
    """Maestros are global and persisted -- 'reiniciar' is for a wrong
    ciclo/ventas upload, not for clearing the territory rules (those get
    replaced by uploading a new maestros workbook, not by reiniciar)."""
    session = {"X-Session-Id": "para-reiniciar"}
    client.post(
        "/api/convatec/tablas-maestras",
        data={"file": (_maestros_xlsx_bytes(), "maestros.xlsx")},
        content_type="multipart/form-data",
        headers=session,
    )
    client.post(
        "/api/convatec/ventas/productos",
        data={"file": (_productos_xlsx_bytes(), "productos.xlsx")},
        content_type="multipart/form-data",
        headers=session,
    )
    reset_resp = client.post("/api/convatec/reiniciar", headers=session)
    assert reset_resp.status_code == 200

    # Ventas were cleared by reiniciar -- nothing to process now.
    procesar_resp = client.post("/api/convatec/procesar", json={"mes": 3}, headers=session)
    assert procesar_resp.status_code == 400
    assert "insumo" in procesar_resp.get_json()["message"]


def test_master_tables_status_reflects_global_upload_from_any_session(client):
    resp_before = client.get("/api/convatec/tablas-maestras", headers={"X-Session-Id": "estado-antes"})
    assert resp_before.get_json()["data"] is None

    client.post(
        "/api/convatec/tablas-maestras",
        data={"file": (_maestros_xlsx_bytes(), "maestros.xlsx")},
        content_type="multipart/form-data",
        headers={"X-Session-Id": "quien-subio"},
    )

    resp_after = client.get("/api/convatec/tablas-maestras", headers={"X-Session-Id": "otra-sesion-cualquiera"})
    body = resp_after.get_json()["data"]
    assert body["representantesRows"] == 1


def test_master_tables_are_global_across_sessions(client):
    """Request 2026-08-06: one Convatec, one set of maestros -- uploading in
    one session must make them usable in another, unlike ventas/resultado."""
    client.post(
        "/api/convatec/tablas-maestras",
        data={"file": (_maestros_xlsx_bytes(), "maestros.xlsx")},
        content_type="multipart/form-data",
        headers={"X-Session-Id": "session-a"},
    )
    client.post(
        "/api/convatec/ventas/productos",
        data={"file": (_productos_xlsx_bytes(), "productos.xlsx")},
        content_type="multipart/form-data",
        headers={"X-Session-Id": "session-b"},
    )
    resp_b = client.post("/api/convatec/procesar", json={"mes": 3}, headers={"X-Session-Id": "session-b"})
    assert resp_b.status_code == 200


def test_ventas_and_resultado_stay_isolated_per_session(client):
    client.post(
        "/api/convatec/tablas-maestras",
        data={"file": (_maestros_xlsx_bytes(), "maestros.xlsx")},
        content_type="multipart/form-data",
        headers={"X-Session-Id": "session-c"},
    )
    client.post(
        "/api/convatec/ventas/productos",
        data={"file": (_productos_xlsx_bytes(), "productos.xlsx")},
        content_type="multipart/form-data",
        headers={"X-Session-Id": "session-c"},
    )
    # session-d shares the global maestros but has no ventas of its own.
    resp_d = client.post("/api/convatec/procesar", json={"mes": 3}, headers={"X-Session-Id": "session-d"})
    assert resp_d.status_code == 400


def test_master_tables_persist_across_a_new_store_instance(client, tmp_path, monkeypatch):
    """The whole point of this feature: a fresh ConvatecSessionStore (e.g.
    after a backend restart) must pick up whatever was last uploaded, from
    disk, with no re-upload."""
    from app.convatec import master_tables_store
    from app.convatec.session_store import ConvatecSessionStore

    monkeypatch.setattr(master_tables_store, "DATA_DIR", tmp_path / "convatec")
    client.post(
        "/api/convatec/tablas-maestras",
        data={"file": (_maestros_xlsx_bytes(), "maestros.xlsx")},
        content_type="multipart/form-data",
        headers={"X-Session-Id": "session-e"},
    )

    fresh_store = ConvatecSessionStore()
    assert fresh_store.has_master_tables("cualquier-sesion-nueva") is True
