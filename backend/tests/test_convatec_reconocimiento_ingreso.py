import io

import openpyxl
import pytest

from app.convatec.reconocimiento_ingreso import ReconocimientoIngresoError, load_total_reconocimiento


class _Upload:
    """Minimal stand-in for a werkzeug FileStorage (see test_excel_loader.py)."""

    def __init__(self, filename, stream):
        self.filename = filename
        self._stream = stream

    def __getattr__(self, name):
        return getattr(self._stream, name)


def _build_reconocimiento_workbook(credit_values):
    """Mimics the real SAP export: metadata rows, then a header row at
    (0-indexed) row 8, then one journal-entry line per row."""
    wb = openpyxl.Workbook()
    ws = wb.active
    headers = ["JE line", "GL account", "Debit amount", "Credit amount", "PK", "Cost center"]
    for i, h in enumerate(headers):
        ws.cell(row=9, column=1 + i, value=h)  # Excel row 9 = 0-indexed row 8
    for r, credit in enumerate(credit_values):
        ws.cell(row=10 + r, column=1, value=r + 1)
        ws.cell(row=10 + r, column=4, value=credit)
        ws.cell(row=10 + r, column=5, value=50)

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return _Upload("reconocimiento.xlsx", buf)


def test_sums_credit_amount_column():
    upload = _build_reconocimiento_workbook([56500.0, 78530.0, 99260.0])
    assert load_total_reconocimiento(upload) == 234290.0


def test_missing_credit_amount_column_raises_user_facing_error():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.cell(row=9, column=1, value="JE line")
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)

    with pytest.raises(ReconocimientoIngresoError, match="Credit amount"):
        load_total_reconocimiento(_Upload("reconocimiento.xlsx", buf))


def test_rejects_non_excel_extension():
    buf = io.BytesIO(b"not an excel file")
    with pytest.raises(ReconocimientoIngresoError, match="Excel"):
        load_total_reconocimiento(_Upload("reconocimiento.csv", buf))
