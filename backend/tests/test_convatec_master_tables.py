import io

import openpyxl
import pytest

from app.convatec.master_tables import MasterTablesError, load_master_tables


class _Upload:
    """Minimal stand-in for a werkzeug FileStorage (see test_excel_loader.py)."""

    def __init__(self, filename, stream):
        self.filename = filename
        self._stream = stream

    def __getattr__(self, name):
        return getattr(self._stream, name)


def _build_tables_workbook():
    """Mimics the real workbook's "Tables" sheet layout: a label two rows
    above the header row, then data -- multiple tables side by side."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Tables"

    # PRODUCTO at column B (index 1); REPRESENTANTES at K (10); CONVENIOS at O (14).
    ws["B2"] = "PRODUCTO"
    ws["K2"] = "REPRESENTANTES"
    ws["O2"] = "CONVENIOS"

    producto_headers = ["SAP Code", "Franquicia", "Familia", "Clase", "Grupo", "Clasificación Cuota Producto"]
    for i, h in enumerate(producto_headers):
        ws.cell(row=4, column=2 + i, value=h)
    producto_rows = [
        ["1004113", "CONTINENCE", "Flexiseal", "0", "Continence Care", "1. CCC"],
        ["1002655", "OSTOMIAS", "GNL", "0", "Ostomy Care", "2. OST"],
    ]
    for r, row in enumerate(producto_rows):
        for c, v in enumerate(row):
            ws.cell(row=5 + r, column=2 + c, value=v)

    rep_headers = ["Convenio", "Grupo de vendedores", "Representante"]
    for i, h in enumerate(rep_headers):
        ws.cell(row=4, column=11 + i, value=h)
    rep_rows = [
        ["SALUD TOTAL", 45, "Katerine Bonilla"],
        ["MEDAXA COLPATRIA SEGUROS SA", None, "Repartir Antioquia"],
    ]
    for r, row in enumerate(rep_rows):
        for c, v in enumerate(row):
            ws.cell(row=5 + r, column=11 + c, value=v)

    convenio_headers = ["CONVENIO", "CIUDAD", "CLIENTE", "CODIGO SAP"]
    for i, h in enumerate(convenio_headers):
        ws.cell(row=4, column=15 + i, value=h)
    ws.cell(row=5, column=15, value="SALUD TOTAL")
    ws.cell(row=5, column=16, value="BOGOTA")
    ws.cell(row=5, column=17, value="SALUD TOTAL EPS")
    ws.cell(row=5, column=18, value="12345")

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return _Upload("maestros.xlsx", buf)


def test_loads_producto_representantes_and_convenios_tables():
    tables = load_master_tables(_build_tables_workbook())

    producto = tables["PRODUCTO"]
    assert list(producto.columns) == [
        "SAP Code",
        "Franquicia",
        "Familia",
        "Clase",
        "Grupo",
        "Clasificación Cuota Producto",
    ]
    assert producto.shape[0] == 2
    assert producto.iloc[0]["Franquicia"] == "CONTINENCE"

    representantes = tables["REPRESENTANTES"]
    assert list(representantes.columns) == ["Convenio", "Grupo de vendedores", "Representante"]
    assert representantes.iloc[1]["Representante"] == "Repartir Antioquia"

    convenios = tables["CONVENIOS"]
    assert list(convenios.columns) == ["CONVENIO", "CIUDAD", "CLIENTE", "CODIGO SAP"]
    assert convenios.iloc[0]["CIUDAD"] == "BOGOTA"


def test_missing_tables_sheet_raises_user_facing_error():
    wb = openpyxl.Workbook()
    wb.active.title = "NotTables"
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)

    with pytest.raises(MasterTablesError, match="hoja llamada 'Tables'"):
        load_master_tables(_Upload("maestros.xlsx", buf))


def test_missing_required_table_raises_user_facing_error():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Tables"
    ws["B2"] = "PRODUCTO"
    ws["B4"] = "SAP Code"
    ws["B5"] = "1"
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)

    with pytest.raises(MasterTablesError, match="REPRESENTANTES"):
        load_master_tables(_Upload("maestros.xlsx", buf))


def test_rejects_non_excel_extension():
    buf = io.BytesIO(b"not an excel file")

    with pytest.raises(MasterTablesError, match="Excel"):
        load_master_tables(_Upload("maestros.csv", buf))


def test_duplicate_table_label_raises_user_facing_error():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Tables"
    ws["B2"] = "PRODUCTO"
    ws["K2"] = "PRODUCTO"  # duplicate label
    ws["B4"] = "SAP Code"
    ws["B5"] = "1"
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)

    with pytest.raises(MasterTablesError, match="más de una vez"):
        load_master_tables(_Upload("maestros.xlsx", buf))
